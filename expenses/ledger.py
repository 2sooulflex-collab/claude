#!/usr/bin/env python3
"""Учет расходов 2Flex2Sooul: данные в CSV -> рабочая таблица xlsx для Google Таблиц.

Команды:
  add [файл.json]   добавить платежи (JSON-массив, как в промте; без файла — из stdin)
  batch "Партия" units=1780 boxes=130 per_box=14 liters=60 shipped=500
                    обновить количества на листе Партии
  check             проверить данные против справочников
  build             собрать out/Расходы 2Flex2Sooul.xlsx
  summary           посчитать Сводку в Python (для сверки с таблицей и ответов в чате)
  import файл.xlsx  забрать Транзакции и Партии из выгрузки Google Таблицы (правки руками, восстановление)
"""
import csv
import datetime as dt
import json
import sys
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
LISTS = json.loads((ROOT / "lists.json").read_text(encoding="utf-8"))
DATA = ROOT / "data"
TX_CSV = DATA / "transactions.csv"
BATCH_CSV = DATA / "batches.csv"
OUT = ROOT / "out" / "Расходы 2Flex2Sooul.xlsx"

TX_COLS = ["Дата", "Категория расхода", "Тип изделия", "Сезон", "Поставщик/производство",
           "Артикул/партия", "Тип отнесения", "Коробок, шт", "Сумма", "Комментарий"]
BATCH_COLS = ["Артикул/партия", "Тип изделия", "Сезон", "Поставщик/производство",
              "Единиц в партии", "Коробок, шт", "Единиц в коробке", "Литраж коробки",
              "Отгружено единиц"]
BATCH_KEYS = {"units": 4, "boxes": 5, "per_box": 6, "liters": 7, "shipped": 8}
# ключи JSON из промта -> колонки листа Транзакции
JSON_KEYS = {"дата": 0, "категория": 1, "тип_изделия": 2, "сезон": 3, "поставщик": 4,
             "партия": 5, "тип_отнесения": 6, "коробок": 7, "сумма": 8, "комментарий": 9}

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
MANUAL_FILL = PatternFill("solid", fgColor="FFFF99")
TOTAL_FILL = PatternFill("solid", fgColor="EDEDED")
BOLD = Font(bold=True)
MONEY = "#,##0"
MONEY2 = "#,##0.00"
TX = "'Транзакции'!"
TX_ROWS = 5000  # сколько строк Транзакций покрыто выпадающими списками


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    return rows[0], rows[1:]


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def batches():
    return read_csv(BATCH_CSV)[1]


def num(v):
    if v in ("", None):
        return None
    f = float(str(v).replace(" ", "").replace(",", "."))
    return int(f) if f.is_integer() else f


def validate(row, batch_names):
    """Список проблем строки транзакции (пустой — все хорошо)."""
    issues = []
    checks = [(1, "categories", "категория"), (2, "product_types", "тип изделия"),
              (3, "seasons", "сезон"), (4, "suppliers", "поставщик"),
              (6, "allocation_types", "тип отнесения")]
    for i, key, label in checks:
        if row[i] and row[i] not in LISTS[key]:
            issues.append(f"{label} «{row[i]}» не из справочника")
    if not row[1]:
        issues.append("пустая категория")
    if row[5] and row[5] not in batch_names:
        issues.append(f"партия «{row[5]}» не из листа Партии")
    try:
        dt.datetime.strptime(row[0], "%d.%m.%Y")
    except ValueError:
        issues.append(f"дата «{row[0]}» не в формате ДД.ММ.ГГГГ")
    try:
        if num(row[8]) is None:
            issues.append("нет суммы")
    except ValueError:
        issues.append(f"сумма «{row[8]}» не число")
    return issues


def warnings(row):
    """Некритичные замечания: строка добавляется, но в себестоимость партий не попадет."""
    if row[6] == "Прямой" and not row[5]:
        return ["прямой расход без партии — в себестоимость не попадет, в Сводке виден как «не разнесено»"]
    return []


def cmd_add(args):
    src = open(args[0], encoding="utf-8") if args else sys.stdin
    items = json.load(src)
    if isinstance(items, dict):
        items = [items]
    today = dt.datetime.now(ZoneInfo("Europe/Moscow")).strftime("%d.%m.%Y")
    header, rows = read_csv(TX_CSV)
    by_batch = {b[0]: b for b in batches()}
    new = []
    for it in items:
        row = [""] * len(TX_COLS)
        for k, v in it.items():
            if k not in JSON_KEYS:
                sys.exit(f"неизвестное поле {k!r}")
            row[JSON_KEYS[k]] = "" if v is None else str(v).strip()
        row[0] = row[0] or today
        row[8] = str(num(row[8]))
        if row[7]:
            row[7] = str(num(row[7]))
        b = by_batch.get(row[5])
        if b:  # тип, сезон и поставщик берем из партии, если не заданы
            for i in (2, 3, 4):
                if not row[i] or (i == 3 and row[i] == "не привязан") or (i == 4 and row[i] == "не указан"):
                    row[i] = b[i - 1]
        row[3] = row[3] or "не привязан"
        row[4] = row[4] or "не указан"
        issues = validate(row, by_batch)
        if issues:
            sys.exit(f"строка не добавлена: {row}\n  " + "\n  ".join(issues))
        for w in warnings(row):
            print(f"внимание: {w}: {row[8]} | {row[9]}")
        new.append(row)
    write_csv(TX_CSV, header, rows + new)
    first = len(rows) + 2
    for i, r in enumerate(new):
        print(f"строка {first + i}: {r[0]} | {r[1]} | {r[5] or '—'} | {r[6]} | {r[8]} | {r[9]}")


def cmd_batch(args):
    name, updates = args[0], args[1:]
    header, rows = read_csv(BATCH_CSV)
    for r in rows:
        if r[0] == name:
            for u in updates:
                k, v = u.split("=", 1)
                r[BATCH_KEYS[k]] = str(num(v))
            write_csv(BATCH_CSV, header, rows)
            print(r)
            return
    sys.exit(f"нет партии {name!r}; есть: {[r[0] for r in rows]}")


def cmd_check(_):
    _, rows = read_csv(TX_CSV)
    names = {b[0] for b in batches()}
    bad = 0
    for i, r in enumerate(rows, start=2):
        for issue in validate(r, names) + warnings(r):
            bad += 1
            print(f"строка {i}: {issue} | {r[8]} | {r[9]}")
    total = sum(num(r[8]) or 0 for r in rows)
    print(f"строк: {len(rows)}, сумма: {total:,.0f}, замечаний: {bad}".replace(",", " "))


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_instructions(wb, nb):
    ws = wb.active
    ws.title = "Инструкция"
    lines = [
        ("Учет расходов 2Flex2Sooul — рабочая таблица", "title"),
        ("", None),
        ("Как вносить расходы", "h"),
        ("Пишите платежи в чат Claude — как раньше писали боту: «#доставка с тяка до лимона 15 мешков куртки осень 5000р». "
         "Можно несколько сумм в одном сообщении и фото платежки. Claude разносит каждую сумму отдельной строкой в Транзакции "
         "и обновляет эту таблицу.", None),
        ("Важно: Claude каждый раз выкладывает обновленную таблицу в папку «2Flex2Sooul — Учет расходов», а прошлую версию "
         "отправляет в корзину. Открывайте таблицу из папки и делитесь папкой, а не файлом.", None),
        ("Правки прямо в таблице (исправить категорию, заполнить Партии) сохранятся: перед обновлением Claude сверяет таблицу "
         "и забирает ваши изменения. Надежнее — написать правку в чат: «строка 21 — общий расход», «приемка куртки осень 1780 шт 130 коробок».", None),
        ("", None),
        ("Листы", "h"),
        ("Транзакции — по строке на каждый платеж. Колонка «Проверка» подсвечивает строки, которые не попадут в себестоимость.", None),
        ("Партии — количества по каждой партии из приемки (желтые ячейки). Без единиц в партии себестоимость на единицу не посчитается.", None),
        ("Сводка — считается формулами: расходы по партиям в разрезе категорий, доля общих расходов, себестоимость на единицу, сверка и итоги по категориям.", None),
        ("Справочники — закрытые списки для выпадающих списков.", None),
        ("", None),
        ("Как считается себестоимость на единицу", "h"),
        ("1) Прямые расходы — строки с типом отнесения «Прямой» и указанной партией. Идут в эту партию целиком.", None),
        ("2) Общие и смешанные расходы с конкретным сезоном (осень, весна, школа, круглый год…) делятся между партиями этого сезона "
         "пропорционально единицам в партии.", None),
        ("3) Общие и смешанные расходы с сезоном «не привязан» (ЧЗ, фулфилмент, подписки, реклама) делятся между всеми партиями "
         "пропорционально единицам.", None),
        ("4) На единицу считается две цифры: на весь объем партии (итог по сезону) и на отгруженные единицы (промежуточная сверка с юниткой).", None),
        ("5) Блок «Сверка» на листе Сводка показывает, сколько денег пока не разнесено по партиям: пока в Партиях не заполнены единицы, "
         "общие расходы висят там.", None),
        ("", None),
        ("Что осталось решить", "h"),
        ("1) Распределение общих и смешанных расходов — сейчас пропорционально единицам, можно переделать на коробки.", None),
        ("2) Какая цифра на единицу основная — на весь объем или на отгруженные.", None),
        ("3) Откуда берутся количества — лист Партии вручную или сообщением в чат при приемке.", None),
        ("4) Размерники и этикетки отнесены к Производственным, раньше такой категории не было — сверить с Полиной.", None),
    ]
    for i, (text, kind) in enumerate(lines, start=2):
        c = ws.cell(i, 2, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            c.font = Font(bold=True, size=14)
        elif kind == "h":
            c.font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 120


def sheet_transactions(wb, rows, nb):
    ws = wb.create_sheet("Транзакции")
    ws.append(TX_COLS + ["Проверка"])
    style_header(ws, 1, len(TX_COLS) + 1)
    for i, r in enumerate(rows, start=2):
        vals = list(r)
        vals[0] = dt.datetime.strptime(r[0], "%d.%m.%Y").date()
        vals[7] = num(r[7])
        vals[8] = num(r[8])
        vals = [None if v == "" else v for v in vals]
        vals.append(f'=IF(I{i}="","⚠ нет суммы",IF(AND(G{i}="Прямой",F{i}=""),"⚠ прямой без партии",""))')
        ws.append(vals)
        ws.cell(i, 1).number_format = "DD.MM.YYYY"
        ws.cell(i, 9).number_format = MONEY
        ws.cell(i, 11).font = Font(color="C00000")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{max(len(rows) + 1, 2)}"
    widths(ws, [11, 30, 22, 13, 22, 30, 14, 11, 13, 50, 22])
    last = TX_ROWS + 1
    lists = [("B", f"'Справочники'!$A$2:$A${1 + len(LISTS['categories'])}"),
             ("C", f"'Справочники'!$C$2:$C${1 + len(LISTS['product_types'])}"),
             ("D", f"'Справочники'!$E$2:$E${1 + len(LISTS['seasons'])}"),
             ("E", f"'Справочники'!$G$2:$G${1 + len(LISTS['suppliers'])}"),
             ("F", f"'Партии'!$A$4:$A${3 + nb}"),
             ("G", f"'Справочники'!$I$2:$I${1 + len(LISTS['allocation_types'])}")]
    for col, ref in lists:
        dv = DataValidation(type="list", formula1=ref, allow_blank=True)
        dv.add(f"{col}2:{col}{last}")
        ws.add_data_validation(dv)


def sheet_batches(wb, brows):
    ws = wb.create_sheet("Партии")
    ws["A1"] = "Справочник партий — количества из приемки"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Желтые колонки заполняются руками (или сообщением в чат). Без единиц в партии себестоимость на единицу не посчитается."
    ws.append(BATCH_COLS)
    style_header(ws, 3, len(BATCH_COLS))
    for r in brows:
        ws.append([None if v == "" else (num(v) if i >= 4 else v) for i, v in enumerate(r)])
        for c in range(5, 10):
            ws.cell(ws.max_row, c).fill = MANUAL_FILL
    ws.freeze_panes = "A4"
    widths(ws, [32, 22, 13, 22, 14, 12, 14, 14, 14])


def sheet_summary(wb, brows):
    ws = wb.create_sheet("Сводка")
    cats = LISTS["categories"]
    nb = len(brows)
    first, last = 4, 3 + nb
    total_row = last + 1
    ws["A1"] = "Сводка расходов и себестоимость на единицу"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Все цифры считаются формулами из листов Транзакции и Партии."
    head = ["Артикул/партия", "Тип изделия", "Сезон", "Единиц в партии", "Отгружено единиц"] + cats + [
        "Итого прямых", "Доля общих и смешанных", "Всего расходов", "На единицу (весь объем)", "На единицу (отгружено)"]
    for c, h in enumerate(head, start=1):
        ws.cell(3, c, h)
    style_header(ws, 3, len(head))
    ws.row_dimensions[3].height = 45
    cat_first = 6
    cat_last = cat_first + len(cats) - 1
    L = get_column_letter
    c_direct, c_share, c_total, c_unit, c_ship = (L(cat_last + k) for k in range(1, 6))
    units = f"$D${first}:$D${last}"
    seasons = f"$C${first}:$C${last}"
    shared = ('(SUMIFS({tx}$I:$I,{tx}$D:$D,{s},{tx}$G:$G,"Общий")'
              '+SUMIFS({tx}$I:$I,{tx}$D:$D,{s},{tx}$G:$G,"Смешанный"))')
    for i, b in enumerate(brows):
        r = first + i
        ws.cell(r, 1, b[0])
        ws.cell(r, 2, b[1])
        ws.cell(r, 3, b[2])
        ws.cell(r, 4, f"=IFERROR(INDEX('Партии'!$E:$E,MATCH($A{r},'Партии'!$A:$A,0))*1,0)")
        ws.cell(r, 5, f"=IFERROR(INDEX('Партии'!$I:$I,MATCH($A{r},'Партии'!$A:$A,0))*1,0)")
        for c in range(cat_first, cat_last + 1):
            col = L(c)
            ws.cell(r, c, f"=SUMIFS({TX}$I:$I,{TX}$F:$F,$A{r},{TX}$B:$B,{col}$3,{TX}$G:$G,\"Прямой\")")
        ws.cell(r, cat_last + 1, f"=SUM({L(cat_first)}{r}:{L(cat_last)}{r})")
        season_pool = shared.format(tx=TX, s=f"$C{r}")
        global_pool = (shared.format(tx=TX, s='"не привязан"') + "+" + shared.format(tx=TX, s='""'))
        ws.cell(r, cat_last + 2,
                f'=IF(OR($C{r}="не привязан",$C{r}=""),0,IFERROR({season_pool}*$D{r}/SUMIF({seasons},$C{r},{units}),0))'
                f"+IFERROR(({global_pool})*$D{r}/SUM({units}),0)")
        ws.cell(r, cat_last + 3, f"={c_direct}{r}+{c_share}{r}")
        ws.cell(r, cat_last + 4, f'=IFERROR({c_total}{r}/$D{r},"—")')
        ws.cell(r, cat_last + 5, f'=IFERROR({c_total}{r}/$E{r},"—")')
        for c in range(4, cat_last + 4):
            ws.cell(r, c).number_format = MONEY
        for c in (cat_last + 4, cat_last + 5):
            ws.cell(r, c).number_format = MONEY2
    ws.cell(total_row, 1, "ИТОГО")
    for c in range(4, cat_last + 4):
        ws.cell(total_row, c, f"=SUM({L(c)}{first}:{L(c)}{last})").number_format = MONEY
    for c in range(1, len(head) + 1):
        ws.cell(total_row, c).font = BOLD
        ws.cell(total_row, c).fill = TOTAL_FILL

    r = total_row + 2
    ws.cell(r, 1, "Сверка").font = Font(bold=True, size=12)
    recon = [
        ("Всего расходов в Транзакциях", f"=SUM({TX}$I:$I)"),
        ("Разнесено по партиям", f"={c_total}{total_row}"),
        ("Не разнесено по партиям", f"=B{r + 1}-B{r + 2}"),
        ("   в т.ч. прямые без партии", f'=SUMIFS({TX}$I:$I,{TX}$G:$G,"Прямой",{TX}$F:$F,"")'),
        ("   в т.ч. общие и смешанные: нет единиц в Партиях или нет партий этого сезона", f"=B{r + 3}-B{r + 4}"),
    ]
    for k, (label, f) in enumerate(recon, start=1):
        ws.cell(r + k, 1, label)
        ws.cell(r + k, 2, f).number_format = MONEY
    ws.cell(r + 3, 1).font = BOLD
    ws.cell(r + 3, 2).font = BOLD

    r = r + len(recon) + 2
    ws.cell(r, 1, "Расходы по категориям").font = Font(bold=True, size=12)
    for c, h in enumerate(["Категория", "Всего", "Прямые", "Общие", "Смешанные"], start=1):
        ws.cell(r + 1, c, h)
    style_header(ws, r + 1, 5)
    for k, cat in enumerate(cats, start=r + 2):
        ws.cell(k, 1, cat)
        ws.cell(k, 2, f"=SUMIFS({TX}$I:$I,{TX}$B:$B,$A{k})")
        for c, t in ((3, "Прямой"), (4, "Общий"), (5, "Смешанный")):
            ws.cell(k, c, f'=SUMIFS({TX}$I:$I,{TX}$B:$B,$A{k},{TX}$G:$G,"{t}")')
        for c in range(2, 6):
            ws.cell(k, c).number_format = MONEY
    k = r + 2 + len(cats)
    ws.cell(k, 1, "ИТОГО")
    for c in range(2, 6):
        ws.cell(k, c, f"=SUM({L(c)}{r + 2}:{L(c)}{k - 1})").number_format = MONEY
    for c in range(1, 6):
        ws.cell(k, c).font = BOLD
        ws.cell(k, c).fill = TOTAL_FILL

    ws.freeze_panes = "B4"
    widths(ws, [32, 20, 13, 11, 11] + [13] * len(cats) + [13, 13, 13, 12, 12])


def sheet_lists(wb):
    ws = wb.create_sheet("Справочники")
    cols = [(1, "Категории расхода", "categories"), (3, "Типы изделий", "product_types"),
            (5, "Сезоны", "seasons"), (7, "Поставщик/производство", "suppliers"),
            (9, "Тип отнесения", "allocation_types")]
    for c, title, key in cols:
        ws.cell(1, c, title)
        for i, v in enumerate(LISTS[key], start=2):
            ws.cell(i, c, v)
        ws.column_dimensions[get_column_letter(c)].width = 38 if c == 1 else 24
        ws.column_dimensions[get_column_letter(c + 1)].width = 3
    for c, *_ in cols:
        ws.cell(1, c).font = BOLD
        ws.cell(1, c).fill = HEADER_FILL


def cmd_build(_):
    _, rows = read_csv(TX_CSV)
    brows = batches()
    wb = Workbook()
    sheet_instructions(wb, len(brows))
    sheet_transactions(wb, rows, len(brows))
    sheet_batches(wb, brows)
    sheet_summary(wb, brows)
    sheet_lists(wb)
    wb.active = 1
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"{OUT} ({OUT.stat().st_size} байт, строк: {len(rows)})")


def cmd_summary(_):
    """Те же расчеты, что формулы листа Сводка, — для сверки с Google Таблицей и ответов в чате."""
    _, rows = read_csv(TX_CSV)
    brows = batches()
    units = {b[0]: num(b[4]) or 0 for b in brows}
    shipped = {b[0]: num(b[8]) or 0 for b in brows}
    season_of = {b[0]: b[2] for b in brows}
    total_units = sum(units.values())
    direct, pool = {}, {}
    for r in rows:
        amount = num(r[8]) or 0
        if r[6] == "Прямой" and r[5] in units:
            direct[r[5]] = direct.get(r[5], 0) + amount
        elif r[6] in ("Общий", "Смешанный"):
            key = r[3] if r[3] not in ("", "не привязан") else None
            pool[key] = pool.get(key, 0) + amount
    fmt = lambda v: f"{v:,.0f}".replace(",", " ")
    placed = 0
    for name in units:
        s = season_of[name]
        share = 0
        if s not in ("", "не привязан"):
            season_units = sum(u for n, u in units.items() if season_of[n] == s)
            share += pool.get(s, 0) * units[name] / season_units if season_units else 0
        share += pool.get(None, 0) * units[name] / total_units if total_units else 0
        total = direct.get(name, 0) + share
        placed += total
        per = lambda q: f"{total / q:,.2f}".replace(",", " ") if q else "—"
        print(f"{name}: прямые {fmt(direct.get(name, 0))}, доля общих {fmt(share)}, всего {fmt(total)}, "
              f"на ед. {per(units[name])}, на отгруж. {per(shipped[name])}")
    grand = sum(num(r[8]) or 0 for r in rows)
    print(f"Всего в Транзакциях {fmt(grand)}, разнесено {fmt(placed)}, не разнесено {fmt(grand - placed)}")
    by_cat = {}
    for r in rows:
        by_cat[r[1]] = by_cat.get(r[1], 0) + (num(r[8]) or 0)
    for c in LISTS["categories"]:
        print(f"  {c}: {fmt(by_cat.get(c, 0))}")


def cmd_import(args):
    from openpyxl import load_workbook
    wb = load_workbook(args[0], data_only=True)

    def cell(v, i):
        if v is None:
            return ""
        if isinstance(v, (dt.datetime, dt.date)):
            return v.strftime("%d.%m.%Y")
        if isinstance(v, float):
            return str(num(v))
        return str(v).strip() if i != 0 else str(v)

    tx = [[cell(v, i) for i, v in enumerate(r[:len(TX_COLS)])]
          for r in wb["Транзакции"].iter_rows(min_row=2, values_only=True)
          if any(v not in (None, "") for v in r[:len(TX_COLS)])]
    bt = [[cell(v, i) for i, v in enumerate(r[:len(BATCH_COLS)])]
          for r in wb["Партии"].iter_rows(min_row=4, values_only=True) if r[0]]
    write_csv(TX_CSV, TX_COLS, tx)
    write_csv(BATCH_CSV, BATCH_COLS, bt)
    print(f"Транзакций: {len(tx)}, партий: {len(bt)}")
    cmd_check([])


if __name__ == "__main__":
    cmds = {"add": cmd_add, "batch": cmd_batch, "check": cmd_check, "build": cmd_build, "summary": cmd_summary, "import": cmd_import}
    if len(sys.argv) < 2 or sys.argv[1] not in cmds:
        sys.exit(__doc__)
    cmds[sys.argv[1]](sys.argv[2:])
