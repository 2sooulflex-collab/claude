#!/usr/bin/env python3
"""Готовит запросы для коннектора Google Sheets: заполнить Google Таблицу тем же содержимым, что и xlsx из ledger.py.

  python3 expenses/gsheet.py build OUT_DIR    — таблица с нуля, пишет в OUT_DIR:
      structure.json    — update_spreadsheet: часовой пояс, листы Транзакции/Партии/Сводка/Справочники (sheetId 0..3)
      <лист>.json       — {"range", "values"} для update_formulas, формулы с «;» (у таблицы локаль ru_RU)
      format_1/2.json   — update_spreadsheet: протяжка формул Сводки (copyPaste), форматы, списки, подсветка
  python3 expenses/gsheet.py rows FROM [TO]   — строки Транзакций FROM..TO (номера строк листа, по умолчанию
                                                до последней) как {"range", "values"} для update_formulas
  python3 expenses/gsheet.py batch "ИМЯ"      — количества партии (колонки E:I листа Партии) так же

В Сводке повторяющиеся формулы передаются один раз (верхняя левая ячейка), остальное протягивается copyPaste —
скрипт заранее проверяет, что протяжка даст ровно те же формулы, что строит ledger.py.
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

from openpyxl.utils import get_column_letter, column_index_from_string

import ledger

SHEET_IDS = {"Транзакции": 0, "Партии": 1, "Сводка": 2, "Справочники": 3}
EPOCH = dt.date(1899, 12, 30)


def to_sheets(formula):
    """Формула из openpyxl → вид для таблицы с локалью ru_RU: «;» вместо «,» вне кавычек, имена листов без кавычек."""
    for name in SHEET_IDS:
        formula = formula.replace(f"'{name}'!", f"{name}!")
    out, quoted = [], False
    for ch in formula:
        if ch == '"':
            quoted = not quoted
        out.append(";" if ch == "," and not quoted else ch)
    return "".join(out)


REF = re.compile(r'(?<![A-Za-zА-Яа-я!$])(\$?)([A-Z]{1,2})(\$?)(\d+)(?![\d(])')


def shift(formula, dr, dc):
    """Сдвинуть относительные ссылки формулы, как это делает протяжка в Google Таблицах."""
    def repl(m):
        cabs, col, rabs, row = m.groups()
        if not cabs:
            col = get_column_letter(column_index_from_string(col) + dc)
        if not rabs:
            row = str(int(row) + dr)
        return f"{cabs}{col}{rabs}{row}"
    parts = formula.split('"')
    return '"'.join(REF.sub(repl, p) if i % 2 == 0 else p for i, p in enumerate(parts))


def cell_value(v):
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        d = v.date() if isinstance(v, dt.datetime) else v
        return (d - EPOCH).days
    if isinstance(v, str):
        if v.startswith("="):
            return to_sheets(v)
        if re.fullmatch(r"[\d\s.,:/+\-]+", v):  # чтобы текст не превратился в число или дату
            return "'" + v
    return v


def grid(ws):
    rows = [[cell_value(c.value) for c in r] for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    return rows


def trim(rows):
    """Убрать хвостовые пустые ячейки строк — меньше передавать."""
    out = []
    for r in rows:
        r = list(r)
        while r and r[-1] is None:
            r.pop()
        out.append(r)
    return out


def a1(r, c):
    return f"{get_column_letter(c)}{r}"


def gr(sheet, r1, c1, r2=None, c2=None):
    """GridRange по 1-based строкам/колонкам включительно; r2=None — до конца листа."""
    g = {"sheetId": SHEET_IDS[sheet], "startRowIndex": r1 - 1, "startColumnIndex": c1 - 1, "endColumnIndex": (c2 or c1)}
    if r2 is not None:
        g["endRowIndex"] = r2
    return g


def rgb(hexstr):
    h = hexstr.lstrip("#")
    return {k: round(int(h[i:i + 2], 16) / 255, 3) for k, i in (("red", 0), ("green", 2), ("blue", 4))}


def fmt(rng, **kw):
    """repeatCell с форматом: bold, italic, size, fg, bg, number, wrap, valign."""
    f, fields = {}, []
    tf = {}
    if "bold" in kw:
        tf["bold"] = kw["bold"]
    if "italic" in kw:
        tf["italic"] = kw["italic"]
    if "size" in kw:
        tf["fontSize"] = kw["size"]
    if "fg" in kw:
        tf["foregroundColor"] = rgb(kw["fg"])
    for k in tf:
        fields.append(f"userEnteredFormat.textFormat.{k}")
    if tf:
        f["textFormat"] = tf
    if "bg" in kw:
        f["backgroundColor"] = rgb(kw["bg"])
        fields.append("userEnteredFormat.backgroundColor")
    if "number" in kw:
        typ = "DATE" if "yy" in kw["number"] else "PERCENT" if "%" in kw["number"] else "NUMBER"
        f["numberFormat"] = {"type": typ, "pattern": kw["number"]}
        fields.append("userEnteredFormat.numberFormat")
    if kw.get("wrap"):
        f["wrapStrategy"] = "WRAP"
        fields.append("userEnteredFormat.wrapStrategy")
    if kw.get("valign"):
        f["verticalAlignment"] = kw["valign"]
        fields.append("userEnteredFormat.verticalAlignment")
    return {"repeatCell": {"range": rng, "cell": {"userEnteredFormat": f}, "fields": ",".join(fields)}}


def widths(sheet, px):
    """Ширины колонок; соседние колонки одной ширины — одним запросом."""
    req, i = [], 0
    while i < len(px):
        j = i
        while j + 1 < len(px) and px[j + 1] == px[i]:
            j += 1
        req.append({"updateDimensionProperties": {"range": {"sheetId": SHEET_IDS[sheet], "dimension": "COLUMNS",
                                                            "startIndex": i, "endIndex": j + 1},
                                                  "properties": {"pixelSize": px[i]}, "fields": "pixelSize"}})
        i = j + 1
    return req


def freeze(sheet, rows, cols=0):
    return {"updateSheetProperties": {"properties": {"sheetId": SHEET_IDS[sheet],
                                                     "gridProperties": {"frozenRowCount": rows, "frozenColumnCount": cols}},
                                      "fields": "gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}}


HEAD = dict(bold=True, bg="#D9E2F3", wrap=True, valign="MIDDLE")
MONEY, MONEY2, DATE = "#,##0", "#,##0.00", "dd.mm.yyyy"


def build(out):
    _, rows = ledger.read_csv(ledger.TX_CSV)
    brows = ledger.batches()
    wb = ledger.make_workbook(rows, brows, instructions=False)
    grids = {name: grid(wb[name]) for name in SHEET_IDS}

    # --- протяжка формул в Сводке ---
    nb = len(brows)
    first, last = 4, 3 + nb
    pool_row, nobatch_row, total_row = last + 1, last + 2, last + 3
    ncat = len(ledger.LISTS["categories"])
    cat_first, cat_last = 6, 5 + ncat
    c_sum, c_share, c_total, c_unit, c_ship = (cat_last + k for k in range(1, 6))
    cats_hdr = next(i for i, r in enumerate(grids["Сводка"], start=1) if r and r[0] == "Категория / статья")
    cat_lines = (cats_hdr + 1, next(i for i, r in enumerate(grids["Сводка"], start=1) if r and r[0] == "ИТОГО") - 1)
    fills = [  # (seed r1, c1, r2, c2) -> destination (r1, c1, r2, c2)
        ((first, 4, first, 5), (first, 4, last, 5)),
        ((first, cat_first, first, cat_first), (first, cat_first, last, cat_last)),
        ((first, c_sum, first, c_sum), (first, c_sum, nobatch_row, c_sum)),
        ((first, c_total, first, c_total), (first, c_total, nobatch_row, c_total)),
        ((first, c_share, first, c_share), (first, c_share, last, c_share)),
        ((first, c_unit, first, c_ship), (first, c_unit, last, c_ship)),
        ((pool_row, cat_first, pool_row, cat_first), (pool_row, cat_first, pool_row, cat_last)),
        ((nobatch_row, cat_first, nobatch_row, cat_first), (nobatch_row, cat_first, nobatch_row, cat_last)),
        ((total_row, 4, total_row, 4), (total_row, 4, total_row, c_total)),
        ((cat_lines[0], 3, cat_lines[0], 3), (cat_lines[0], 3, cat_lines[1], 3)),
    ]
    g = grids["Сводка"]
    requests_fill = []
    for (sr1, sc1, sr2, sc2), (dr1, dc1, dr2, dc2) in fills:
        h, w = sr2 - sr1 + 1, sc2 - sc1 + 1
        for r in range(dr1, dr2 + 1):
            for c in range(dc1, dc2 + 1):
                seed = g[sr1 - 1 + (r - dr1) % h][sc1 - 1 + (c - dc1) % w]
                want = g[r - 1][c - 1]
                got = shift(seed, r - (sr1 + (r - dr1) % h), c - (sc1 + (c - dc1) % w))
                assert got == want, f"{a1(r, c)}: протяжка даст {got}, а нужно {want}"
                if not (sr1 <= r <= sr2 and sc1 <= c <= sc2):  # все, кроме самой исходной ячейки, придет протяжкой
                    g[r - 1][c - 1] = None
        requests_fill.append({"copyPaste": {"source": gr("Сводка", sr1, sc1, sr2, sc2),
                                            "destination": gr("Сводка", dr1, dc1, dr2, dc2),
                                            "pasteType": "PASTE_FORMULA"}})

    out.mkdir(parents=True, exist_ok=True)
    structure = [
        {"updateSpreadsheetProperties": {"properties": {"timeZone": "Europe/Moscow"}, "fields": "timeZone"}},
        {"updateSheetProperties": {"properties": {"sheetId": 0, "title": "Транзакции"}, "fields": "title"}},
    ] + [{"addSheet": {"properties": {"sheetId": i, "title": n}}} for n, i in SHEET_IDS.items() if i]
    (out / "structure.json").write_text(json.dumps(structure, ensure_ascii=False))
    for name, rows_ in grids.items():
        rows_ = trim(rows_)
        if name == "Сводка":  # по блокам, пропуская пустые строки между ними
            blocks, start = [], None
            for i, r in enumerate(rows_ + [[]], start=1):
                if r and start is None:
                    start = i
                elif not r and start is not None:
                    blocks.append((start, rows_[start - 1:i - 1]))
                    start = None
            for n, (r0, part) in enumerate(blocks, start=1):
                (out / f"{name}_{n}.json").write_text(
                    json.dumps({"range": f"{name}!A{r0}", "values": part}, ensure_ascii=False))
        else:
            (out / f"{name}.json").write_text(json.dumps({"range": f"{name}!A1", "values": rows_}, ensure_ascii=False))

    # --- оформление ---
    T, P, S, R = "Транзакции", "Партии", "Сводка", "Справочники"
    L = {k: ledger.col(k) for k in ledger.C}
    req = list(requests_fill)
    ntx = len(ledger.TX_COLS)
    req += [fmt(gr(T, 1, 1, 1, ntx), **HEAD), fmt(gr(T, 2, 1, None, 1), number=DATE),
            fmt(gr(T, 2, 10, None, 10), number=MONEY), freeze(T, 1),
            {"setBasicFilter": {"filter": {"range": gr(T, 1, 1, None, ntx)}}}]
    req += widths(T, [90, 230, 150, 150, 95, 160, 230, 110, 75, 95, 460])
    lists = {"cat": "$A$2:$A$10", "art": f"$C$2:$C${1 + len(ledger.ALL_ARTICLES)}", "ptype": "$E$2:$E$11",
             "season": "$G$2:$G$8", "supplier": "$I$2:$I$6", "alloc": "$K$2:$K$4"}
    for key, ref in list(lists.items()) + [("batch", None)]:
        src = f"=Партии!$A$4:$A${3 + nb}" if key == "batch" else f"=Справочники!{ref}"
        c = ledger.C[key] + 1
        req.append({"setDataValidation": {"range": gr(T, 2, c, None, c), "rule": {
            "condition": {"type": "ONE_OF_RANGE", "values": [{"userEnteredValue": src}]},
            "showCustomUi": True, "strict": False}}})
    red = (f'=OR(AND(${L["amount"]}2="";${L["cat"]}2<>"");AND(${L["cat"]}2<>"";${L["alloc"]}2="");'
           f'AND(${L["alloc"]}2="Прямой";${L["batch"]}2=""))')
    with_art = ";".join(f'${L["cat"]}2="{c}"' for c in ledger.ARTICLES)
    orange = f'=AND(OR({with_art});${L["art"]}2="")'
    for i, (formula, color) in enumerate([(red, "#F4CCCC"), (orange, "#FCE5CD")]):
        req.append({"addConditionalFormatRule": {"index": i, "rule": {
            "ranges": [gr(T, 2, 1, None, ntx)],
            "booleanRule": {"condition": {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": formula}]},
                            "format": {"backgroundColor": rgb(color)}}}}})

    req += [fmt(gr(P, 1, 1), bold=True, size=13), fmt(gr(P, 3, 1, 3, 9), **HEAD),
            fmt(gr(P, 4, 5, 3 + nb, 9), bg="#FFF2CC", number=MONEY), freeze(P, 3)]
    req += widths(P, [250, 170, 100, 170, 110, 90, 110, 110, 120])

    sg = grids[S]
    recon_row = total_row + 2
    seas_hdr = recon_row + 8
    seas_last = seas_hdr + len(ledger.LISTS["seasons"])
    cat_total = cat_lines[1] + 1
    req += [fmt(gr(S, 1, 1), bold=True, size=13), fmt(gr(S, 3, 1, 3, c_ship), **HEAD),
            fmt(gr(S, first, 4, total_row, c_total), number=MONEY),
            fmt(gr(S, first, c_unit, last, c_ship), number=MONEY2),
            fmt(gr(S, pool_row, 1, nobatch_row, c_ship), bg="#FCE4D6"),
            fmt(gr(S, total_row, 1, total_row, c_ship), bold=True, bg="#EDEDED"),
            fmt(gr(S, recon_row, 1), bold=True, size=12), fmt(gr(S, recon_row + 1, 1, recon_row + 1, 2), bold=True),
            fmt(gr(S, recon_row + 1, 2, recon_row + 5, 2), number=MONEY),
            fmt(gr(S, seas_hdr - 1, 1), bold=True, size=12), fmt(gr(S, seas_hdr, 1, seas_hdr, 3), **HEAD),
            fmt(gr(S, seas_hdr + 1, 2, seas_last, 3), number=MONEY),
            fmt(gr(S, cats_hdr - 1, 1), bold=True, size=12), fmt(gr(S, cats_hdr, 1, cats_hdr, 5), **HEAD),
            fmt(gr(S, cats_hdr + 1, 2, cat_total, 2), number=MONEY),
            fmt(gr(S, cats_hdr + 1, 3, cat_total, 3), number="0.0%"),
            fmt(gr(S, cats_hdr + 1, 5, cat_total, 5), number=DATE),
            fmt(gr(S, cat_total, 1, cat_total, 5), bold=True, bg="#EDEDED"),
            freeze(S, 3, 1)]
    indented = [r for r in range(cats_hdr + 1, cat_total) if sg[r - 1][0].startswith(" ")]
    while indented:  # подряд идущие строки статей — одним запросом
        r0 = r1 = indented.pop(0)
        while indented and indented[0] == r1 + 1:
            r1 = indented.pop(0)
        req.append(fmt(gr(S, r0, 1, r1, 5), italic=True, fg="#595959"))
    req += widths(S, [300, 150, 100, 95, 95] + [110] * ncat + [110, 110, 110, 100, 100])
    req.append({"updateDimensionProperties": {"range": {"sheetId": SHEET_IDS[S], "dimension": "ROWS",
                                                        "startIndex": 2, "endIndex": 3},
                                              "properties": {"pixelSize": 60}, "fields": "pixelSize"}})
    req += [fmt(gr(R, 1, 1, 1, 11), bold=True, bg="#D9E2F3")] + widths(R, [270, 20, 200, 20, 170, 20, 110, 20, 160, 20, 110])
    half = len(req) // 2
    (out / "format_1.json").write_text(json.dumps(req[:half], ensure_ascii=False))
    (out / "format_2.json").write_text(json.dumps(req[half:], ensure_ascii=False))

    for f in sorted(out.iterdir()):
        print(f.name, f.stat().st_size)


def tx_rows(first, last=None):
    """Строки листа Транзакции first..last из transactions.csv — в том виде, как их пишет build."""
    _, rows = ledger.read_csv(ledger.TX_CSV)
    last = last or len(rows) + 1
    if not 2 <= first <= last <= len(rows) + 1:
        sys.exit(f"в transactions.csv строки листа 2..{len(rows) + 1}")
    out = []
    for r in rows[first - 2:last - 1]:
        vals = list(r)
        vals[ledger.C["date"]] = dt.datetime.strptime(r[ledger.C["date"]], "%d.%m.%Y").date()
        vals[ledger.C["boxes"]] = ledger.num(r[ledger.C["boxes"]])
        vals[ledger.C["amount"]] = ledger.num(r[ledger.C["amount"]])
        out.append(["" if v in ("", None) else cell_value(v) for v in vals])
    end = get_column_letter(len(ledger.TX_COLS))
    return {"range": f"Транзакции!A{first}:{end}{last}", "values": out}


def batch_row(name):
    """Количества партии: колонки E:I ее строки на листе Партии (партии идут с 4-й строки в порядке batches.csv)."""
    for i, b in enumerate(ledger.batches(), start=4):
        if b[0] == name:
            return {"range": f"Партии!E{i}:I{i}", "values": [["" if v == "" else ledger.num(v) for v in b[4:9]]]}
    sys.exit(f"нет партии {name!r}")


if __name__ == "__main__":
    cmd, args = sys.argv[1], sys.argv[2:]
    if cmd == "build":
        build(Path(args[0]))
    elif cmd == "rows":
        print(json.dumps(tx_rows(*map(int, args)), ensure_ascii=False))
    elif cmd == "batch":
        print(json.dumps(batch_row(args[0]), ensure_ascii=False))
    else:
        sys.exit(__doc__)
